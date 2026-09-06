#!/usr/bin/env python3
"""
Build conferences_2026-09-06.tsv and .xlsx for AG Schuck quantum-photonics
conference tracker, incrementally from the 2026-08-30 snapshot.
"""
import csv, os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

TODAY = date(2026, 9, 6)
SNAPSHOT = "2026-09-06"
OUTDIR = "/home/user/conferences"
PREV_TSV = os.path.join(OUTDIR, "conferences_2026-08-30.tsv")

COLS = [
    "id", "name", "acronym", "type", "topic_tags",
    "start_date", "end_date", "location",
    "abstract_deadline", "paper_deadline",
    "registration_deadline", "early_bird_deadline",
    "website", "relevance_score", "relevance_note",
    "status", "last_verified", "last_change", "source_urls",
]
DEADLINE_COLS = ["abstract_deadline", "paper_deadline",
                 "registration_deadline", "early_bird_deadline"]

# ── Load previous snapshot ──────────────────────────────────────────────────
with open(PREV_TSV, newline="", encoding="utf-8") as f:
    reader = csv.DictReader(f, delimiter="\t")
    rows = list(reader)

by_id = {r["id"]: r for r in rows}

def update(id_, **kwargs):
    r = by_id[id_]
    r.update(kwargs)

# ── Row-level updates for this run (2026-09-06) ─────────────────────────────

# 1. Events that concluded between 2026-08-30 and 2026-09-06
update("tqc-2026",
       status="archived",
       last_verified=SNAPSHOT,
       last_change="event concluded; archived (was 2026-08-31 to 2026-09-04)")

# 2. Event starting today
update("asc-2026",
       status="ongoing",
       last_verified=SNAPSHOT,
       last_change="event now in progress (2026-09-06 to 2026-09-11)")

# 3. Bug fix: IEEE QPAIN 2026 concluded back in April 2026 but was left as
#    "upcoming" in prior snapshots — correct it now.
update("ieee-qpain-2026",
       status="archived",
       last_verified=SNAPSHOT,
       last_change="correction: event concluded 2026-04-16/18 but was left as upcoming in prior snapshots; archived")

# 4. EQTC 2026 — official site (eqtc2026.eu) gives corrected start date and
#    registration tiers not previously captured
update("eqtc-2026",
       start_date="2026-11-30",
       registration_deadline="2026-11-02",
       early_bird_deadline="2026-09-14",
       last_verified=SNAPSHOT,
       last_change="start date corrected to Nov 30 (was Nov 29); early-bird registration deadline 2026-09-14 and regular registration deadline 2026-11-02 added per official site",
       source_urls="https://qt.eu/events/eqtc-2026-european-quantum-technologies-conference; https://www.eqtc2026.eu/")

# 5. HDQS 2026 (Benasque) — aggregator sources show materially different
#    dates than previously recorded; official benasque.org page could not be
#    fetched this run (egress blocked), so flagged for confirmation.
update("hdqs-2026",
       start_date="2026-09-20",
       end_date="2026-10-03",
       last_verified=SNAPSHOT,
       last_change="dates revised to Sep 20-Oct 3 (was Oct 4-9) per secondary aggregator; official benasque.org page unreachable this run (network egress blocked) -- recommend manual confirmation next run",
       source_urls="https://benasque.org/2026hdqs/ (unreachable this run); aggregator: quantum.technology/conf/index.html")

# 6. Photonics West 2027 — abstract deadline (2026-07-22) has now passed
update("photonics-west-2027",
       status="cfp_closed",
       last_verified=SNAPSHOT,
       last_change="abstract deadline 2026-07-22 has passed; status now cfp_closed")

# ── New rows discovered this run ────────────────────────────────────────────
def new_row(id_, name, acronym, typ, tags, sd, ed, loc,
            abs_dl, pap_dl, reg_dl, eb_dl, url, score, note, status, change, srcs):
    return {
        "id": id_, "name": name, "acronym": acronym, "type": typ,
        "topic_tags": tags, "start_date": sd, "end_date": ed, "location": loc,
        "abstract_deadline": abs_dl, "paper_deadline": pap_dl,
        "registration_deadline": reg_dl, "early_bird_deadline": eb_dl,
        "website": url, "relevance_score": score, "relevance_note": note,
        "status": status, "last_verified": SNAPSHOT, "last_change": change,
        "source_urls": srcs,
    }

new_rows = [
    new_row("qps-2026",
        "Quantum Photonics Spotlight 2026", "QPS 2026", "Conference",
        "integrated quantum photonics;quantum optics;single-photon sources",
        "2026-09-28", "2026-10-01", "Paderborn, Germany",
        "TBA", "n/a", "TBA", "n/a",
        "https://phoqs.uni-paderborn.de/en/veranstaltungen/quantum-photonics-spotlight-2026",
        5, "PhoQS Paderborn's flagship quantum photonics meeting -- closest major event to Münster, direct topical overlap",
        "upcoming", "new entry; found via web search this run",
        "https://phoqs.uni-paderborn.de/en/veranstaltungen/quantum-photonics-spotlight-2026; https://express.converia.de/frontend/index.php?sub=2189"),

    new_row("erice-qts-2026",
        "Quantum Technology School Erice 2026", "Erice QTS 2026", "Summer School",
        "quantum optics;single-photon sources;integrated quantum photonics",
        "2026-06-23", "2026-06-28", "Erice, Italy",
        "TBA", "n/a", "TBA", "n/a",
        "https://sites.google.com/view/quantumtechnologyschoolerice/home",
        4, "PhD school on quantum photonics/quantum-dot single-photon sources (organizers incl. Trotta, Portalupi) -- event concluded, added for series continuity",
        "archived", "new entry; event already concluded (Jun 23-28 2026), added retroactively so 2027 edition can be tracked",
        "https://sites.google.com/view/quantumtechnologyschoolerice/home"),

    new_row("erice-qts-2027",
        "Quantum Technology School Erice 2027", "Erice QTS 2027", "Summer School",
        "quantum optics;single-photon sources;integrated quantum photonics",
        "TBA", "TBA", "Erice, Italy",
        "TBA", "n/a", "TBA", "n/a",
        "https://sites.google.com/view/quantumtechnologyschoolerice/home",
        4, "PhD school on quantum photonics/quantum-dot single-photon sources -- annual series, PhD-student fit",
        "upcoming", "placeholder; year roll-over from 2026 edition, 2027 dates not yet announced",
        "https://sites.google.com/view/quantumtechnologyschoolerice/home"),

    new_row("ieee-qpain-2027",
        "IEEE 3rd International Conference on Quantum Photonics, AI & Networking", "IEEE QPAIN 2027", "Conference",
        "integrated quantum photonics;quantum networking",
        "2027-04-08", "2027-04-10", "Chattogram, Bangladesh",
        "2026-11-30", "TBA", "TBA", "TBA",
        "https://qpain.org/",
        3, "Quantum photonics and networking conference (IEEE-sponsored); successor to QPAIN 2026",
        "cfp_open", "new entry; year roll-over from 2026 edition (Apr 16-18 2026, now archived)",
        "https://qpain.org/; https://qpain.org/call-for-papers"),

    new_row("iqt-nordics-2027",
        "IQT Nordics 2027", "IQT Nordics 2027", "Conference",
        "quantum communication;quantum computing hardware;integrated quantum photonics",
        "2027-06-07", "2027-06-09", "Copenhagen, Denmark",
        "TBA", "n/a", "TBA", "TBA",
        "https://iqtevent.com/nordics/",
        3, "Nordic quantum industry conference with hardware and QKD sessions; successor to IQT Nordics 2026",
        "upcoming", "new entry; year roll-over from 2026 edition (Oslo), 2027 moves to Copenhagen",
        "https://www.insidequantumtechnology.com/news-archive/announcing-the-return-of-iqt-nordics-in-copenhagen-2027/"),

    new_row("quantum-networks-summit-2027",
        "Quantum Networks Summit 2027", "QNS 2027", "Conference",
        "quantum communication;quantum networking",
        "2027-04-07", "2027-04-08", "Paris, France",
        "TBA", "n/a", "TBA", "TBA",
        "TBA",
        3, "Industry/research summit on quantum networking and QKD infrastructure, co-located with Upperside World Congress",
        "upcoming", "new entry; found via web search this run",
        "unofficial aggregator reports, pending confirmation on an official conference site"),
]

for nr in new_rows:
    by_id[nr["id"]] = nr
    rows.append(nr)

# ── Mark last_verified for a couple more rows we actively re-checked but did
#    not need to change (kept for traceability) ─────────────────────────────
for id_, note in [
    ("cleo-2027", "verified, no change (May 2-7 2027, Long Beach CA confirmed)"),
    ("benasque-qnp-2027", "verified, no change (Mar 28-Apr 2 2027 confirmed)"),
    ("photonics-days-2026", "verified, no change (Oct 7-8 2026 confirmed, partner countries Korea/Taiwan)"),
]:
    r = by_id.get(id_)
    if r:
        r["last_verified"] = SNAPSHOT
        r["last_change"] = note

events = rows

# ── write TSV ────────────────────────────────────────────────────────────────
def write_tsv(events, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLS, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        for e in events:
            w.writerow(e)
    print(f"  TSV  -> {path}  ({len(events)} rows)")

# ── write XLSX ───────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
ARCHIVED_FILL = PatternFill(fill_type="solid", fgColor="E8E8E8")
SCORE_COLORS = {5: "D4EDDA", 4: "D1ECF1", 3: "FFF3CD", 2: "F8D7DA", 1: "F5C6CB"}
RED = PatternFill(fill_type="solid", fgColor="FF4444")
AMBER = PatternFill(fill_type="solid", fgColor="FFAA00")

def parse_date(s):
    if not s or s in ("TBA", "n/a"):
        return None
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None

def days_from_today(s):
    d = parse_date(s)
    if d is None:
        return None
    return (d - TODAY).days

def deadline_fill(s):
    n = days_from_today(s)
    if n is None:
        return None
    if 0 <= n <= 14:
        return RED
    if 15 <= n <= 60:
        return AMBER
    return None

def write_xlsx(events, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Conferences"

    ws.append(COLS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"

    col_idx = {c: i + 1 for i, c in enumerate(COLS)}

    for ev in events:
        row_vals = [ev.get(c, "") for c in COLS]
        ws.append(row_vals)
        row_num = ws.max_row

        is_archived = ev["status"] == "archived"
        if is_archived:
            for cell in ws[row_num]:
                cell.fill = ARCHIVED_FILL

        score_cell = ws.cell(row=row_num, column=col_idx["relevance_score"])
        sc = ev.get("relevance_score")
        if sc and str(sc).isdigit() and int(sc) in SCORE_COLORS:
            score_cell.fill = PatternFill(fill_type="solid", fgColor=SCORE_COLORS[int(sc)])

        for dcol in DEADLINE_COLS:
            val = ev.get(dcol, "")
            fill = deadline_fill(str(val))
            if fill:
                ws.cell(row=row_num, column=col_idx[dcol]).fill = fill

    CAP = {
        "id": 28, "name": 55, "acronym": 18, "type": 16, "topic_tags": 55,
        "start_date": 13, "end_date": 13, "location": 28,
        "abstract_deadline": 17, "paper_deadline": 14,
        "registration_deadline": 20, "early_bird_deadline": 18,
        "website": 48, "relevance_score": 9, "relevance_note": 52,
        "status": 14, "last_verified": 14, "last_change": 50, "source_urls": 60,
    }
    for col_name, cap_w in CAP.items():
        ci = col_idx[col_name]
        ws.column_dimensions[get_column_letter(ci)].width = cap_w

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=False, vertical="top")

    wb.save(path)
    print(f"  XLSX -> {path}")


if __name__ == "__main__":
    tsv_path = os.path.join(OUTDIR, f"conferences_{SNAPSHOT}.tsv")
    xlsx_path = os.path.join(OUTDIR, f"conferences_{SNAPSHOT}.xlsx")

    write_tsv(events, tsv_path)
    write_xlsx(events, xlsx_path)
    print(f"\nDone. {len(events)} events written.")

    from collections import Counter
    stats = Counter(e["status"] for e in events)
    for s, n in sorted(stats.items()):
        print(f"  {s}: {n}")
