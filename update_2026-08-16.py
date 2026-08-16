#!/usr/bin/env python3
"""
Update script for the 2026-08-16 snapshot.
Reads conferences_2026-05-13.tsv, applies archiving of passed events,
research-verified field corrections, TBA-placeholder resolutions, and
newly discovered rows, then writes conferences_2026-08-16.tsv and .xlsx.
"""

# ---- Part 1: build TSV ----
import csv

TODAY = "2026-08-16"
SRC = "/home/user/conferences/conferences_2026-05-13.tsv"
OUT_TSV = "/home/user/conferences/conferences_2026-08-16.tsv"

COLS = [
    "id", "name", "acronym", "type", "topic_tags",
    "start_date", "end_date", "location",
    "abstract_deadline", "paper_deadline",
    "registration_deadline", "early_bird_deadline",
    "website", "relevance_score", "relevance_note",
    "status", "last_verified", "last_change", "source_urls",
]

with open(SRC, newline="", encoding="utf-8") as f:
    reader = csv.DictReader(f, delimiter="\t")
    rows = [dict(r) for r in reader]

by_id = {r["id"]: r for r in rows}

PASSED_IDS = [
    "cleo-2026", "damop-2026", "photonics-north-2026", "cargese-qt-2026",
    "spie-pfq-2026", "islc-2026", "grc-quantum-science-2026", "icap-2026",
    "ecio-2026", "optica-quantum-2026", "icop-2026", "ccw-london-2026",
    "egas-2026", "benasque-qsi-2026", "les-houches-pa-2026", "iqt-nordics-2026",
    "quantum-tech-world-2026", "oecc-2026", "quantum-korea-2026", "spw-2026",
    "icton-2026", "ieee-sum-2026", "grc-plasmonics-nano-2026", "cewqo-2026",
    "apc-2026", "grc-mech-quant-2026", "lake-como-ufqp-2026", "ieee-rapid-2026",
]
for pid in PASSED_IDS:
    r = by_id[pid]
    r["status"] = "archived"
    r["last_verified"] = TODAY
    r["last_change"] = "archived - event has passed"

# ----- near-term (Aug-Dec 2026) verified updates -----
NEAR_TERM = {
    "spie-op-2026": dict(early_bird_deadline="2026-08-07", status="cfp_closed",
        last_change="Early-bird deadline passed Aug 7 (fees +$170 after); post-deadline abstracts still accepted"),
    "qcrypt-2026": dict(abstract_deadline="2026-03-13", paper_deadline="2026-05-01", status="cfp_closed",
        last_change="Talk deadline was 2026-03-13, poster deadline 2026-05-01 (both closed); venue: Learning Crossroads (CRX), University of Ottawa; refund cutoff passed",
        source_urls="https://qcrypt.net/2026/; https://qcrypt.net/2026/attend/registration/"),
    "eosam-2026": dict(abstract_deadline="2026-04-14", early_bird_deadline="2026-06-15", status="cfp_closed",
        last_change="Abstract deadline 2026-04-14 and early-bird 2026-06-15 both passed; standard registration rate now in effect"),
    "tqc-2026": dict(abstract_deadline="2026-04-20", early_bird_deadline="2026-06-28", status="cfp_closed",
        last_change="Poster deadline 2026-04-20 passed; registration early-bird 2026-06-28 passed; accepted-papers list published"),
    "asc-2026": dict(status="cfp_closed",
        last_change="Abstract deadline has passed (exact date unconfirmed, ~Jan/Feb 2026); abstract-change cutoff 2026-08-03 also passed"),
    "inphomir-school-2026": dict(
        last_change="No change found; application deadline still not published, check official site directly"),
    "ieee-qce26": dict(paper_deadline="2026-04-27", status="cfp_closed",
        last_change="Technical paper deadline extended to 2026-04-27 (passed); registration opened 2026-04-14"),
    "ecoc-2026": dict(
        last_change="No change; registration open, fees EUR450-1280"),
    "mne-2026": dict(early_bird_deadline="2026-07-24", status="cfp_closed",
        last_change="Early-bird deadline 2026-07-24 passed, standard rate now in effect"),
    "fio-2026": dict(end_date="2026-10-01", early_bird_deadline="2026-08-14", status="cfp_closed",
        last_change="End date extended to Oct 1 (from Sep 30, now a 4-day Technical Conference); early-bird deadline 2026-08-14 passed; cancellation deadline 2026-09-13"),
    "hdqs-2026": dict(
        last_change="Dates possibly shifted to Oct 5-8 per a secondary source; unconfirmed via official benasque.org (site unreachable this run) - verify next run"),
    "photonics-days-2026": dict(
        last_change="Registration now open; Photonics Travel Days added Oct 6 & 9 flanking main event"),
    "ipc-2026": dict(paper_deadline="2026-05-04", early_bird_deadline="2026-10-08", status="cfp_closed",
        last_change="Paper deadline 2026-05-04 passed; early-bird registration deadline 2026-10-08 (not yet passed); registration now open"),
    "eqtc-2026": dict(start_date="2026-11-30", abstract_deadline="2026-08-03",
        website="https://www.eqtc2026.eu/", status="cfp_closed",
        last_change="Start date corrected to Nov 30 (from Nov 29); abstract deadline 2026-08-03 passed; registration opened Jul 2; canonical URL updated to eqtc2026.eu",
        source_urls="https://www.eqtc2026.eu/; https://qt.eu/news/2026/2026-07-02_EQTC-2026-registration-open-dublin"),
}

# ----- 2027 events verified -----
FUTURE_2027 = {
    "photonics-west-2027": dict(status="cfp_closed",
        last_change="Abstract deadline 2026-07-22 confirmed and has now passed; post-deadline submissions being accepted"),
    "qip-2027": dict(location="Singapore (NUS University Cultural Centre)",
        last_change="Venue refined to NUS University Cultural Centre; program: tutorials Feb 20-21, main program Feb 22-26"),
    "dpg-samop-2027": dict(website="https://www.dpg-physik.de/veranstaltungen/2027/dpg-fruehjahrstagung-hannover",
        last_change="Confirmed Hannover venue (Leibniz University); abstract/registration portal opens Sept 2026 (not yet open)"),
    "ofc-2027": dict(last_change="No change; dates and LA Convention Center venue confirmed"),
    "spie-oo-2027": dict(location="Prague, Czech Republic",
        website="https://spie.org/conferences-and-exhibitions/optics-and-optoelectronics",
        last_change="Venue confirmed: Clarion Congress Hotel, Prague"),
    "aps-summit-2027": dict(last_change="No change; Atlanta Apr 11-16 confirmed"),
    "cleo-europe-2027": dict(location="Munich, Germany",
        last_change="Confirmed as part of World of Photonics Congress 2027 (ICM - International Congress Center); exhibition halls Jun 22-25"),
    "ieee-rapid-2027": dict(
        last_change="Unable to independently verify this run (official site unreachable); some secondary sources suggest possible September dates - flagged for manual check next run"),
    "ipc-2027": dict(
        last_change="Discrepancy found: one aggregator lists Nov 7-11 vs our Nov 10-14; neither confirmed via official ieee-ipc.org (unreachable this run) - flagged for manual verification; Eindhoven venue still tentative"),
}

# ----- TBA placeholder resolutions -----
TBA_RESOLVE = {
    "photonics-switching-2026": dict(
        name="Photonics in Switching and Computing 2026", acronym="PSC 2026",
        start_date="2026-09-15", end_date="2026-09-18", location="Valencia, Spain",
        paper_deadline="2026-06-19", website="https://psc2026.org/", status="cfp_closed",
        last_change="Resolved: PSC 2026, Sep 15-18 Valencia; paper deadline extended to Jun 19 (passed); post-deadline papers accepted through Sep 4",
        source_urls="https://psc2026.org/"),
    "it-fab-school-2026": dict(
        start_date="2026-10-27", end_date="2026-10-30", location="Turin, Italy",
        website="https://www.inrim.it/it/eventi/it-fab-hands-school-micro-and-nano-fabrication",
        status="upcoming",
        last_change="Resolved: first traveling edition hosted at INRiM Turin (PiQuET infrastructure), Oct 27-30; no fixed application deadline, contact itfabschool@it-fab.it",
        source_urls="https://www.inrim.it/it/eventi/it-fab-hands-school-micro-and-nano-fabrication"),
    "opic": dict(
        start_date="2026-04-20", end_date="2026-04-24", location="Yokohama, Japan",
        paper_deadline="2026-01-05", website="https://opicon.jp/", status="archived",
        last_change="Resolved: OPIC2026 was Apr 20-24 2026 Pacifico Yokohama (already passed); next edition (~OPIC2027) not yet announced",
        source_urls="https://opicon.jp/"),
    "owtnm": dict(
        id="owtnm-2026",
        name="32nd Optical Wave & Waveguide Theory and Modelling Workshop",
        acronym="OWTNM 2026 (XXXII)",
        start_date="2026-04-08", end_date="2026-04-10", location="Lausanne, Switzerland",
        abstract_deadline="2026-02-06", registration_deadline="2026-03-06",
        website="https://owtnm26.epfl.ch/", status="archived",
        last_change="Resolved: XXXII OWTNM was Apr 8-10 2026 EPFL Lausanne (already passed); annual (not biennial) cadence; next (XXXIII) 2027 host TBA",
        source_urls="https://owtnm26.epfl.ch/; https://www.owtnm.eu/"),
    "opon-2025": dict(
        start_date="2025-02-12", end_date="2025-02-14", location="Münster, Germany",
        website="https://www.uni-muenster.de/Physik.FT/en/opon2025/index.html", status="archived",
        last_change="Resolved: 8th OPON (Optical Properties of Nanostructures) workshop held Feb 12-14 2025 Münster (AG Kuhn), invitation-only; rotates Wrocław/Münster/Warsaw/Bayreuth; no 9th edition announced yet",
        source_urls="https://www.uni-muenster.de/Physik.FT/en/opon2025/index.html"),
    "nrw-nano-2026": dict(
        id="nrw-nano-2025",
        name="11th NRW NanoConference 2025", acronym="NRW Nano 2025",
        start_date="2025-09-30", end_date="2025-10-01", location="Dortmund, Germany",
        registration_deadline="2025-05-15", website="https://www.nanoconference.de/",
        status="archived",
        last_change="Corrected: actual 11th edition was Sep 30-Oct 1 2025 Dortmund (previously mis-tracked as 2026); biennial cadence, 12th edition (~2027) not yet announced",
        source_urls="https://www.nanoconference.de/"),
    "icqe": dict(
        id="icqe-2025",
        name="International Conference on Quantum Energy 2025",
        start_date="2025-06-03", end_date="2025-06-06", location="Padua, Italy",
        website="https://icqe.com.au/", status="archived",
        last_change="Resolved: real recurring conference (Univ. Padua / Quantum Energy Initiative); last known edition Jun 3-6 2025 Padua; no 2026/2027 edition announced yet",
        source_urls="https://icqe.com.au/"),
}

def apply_updates(d, updates):
    for k, v in updates.items():
        if k == "id":
            continue
        d[k] = v
    d["last_verified"] = TODAY

for uid, upd in NEAR_TERM.items():
    apply_updates(by_id[uid], upd)
for uid, upd in FUTURE_2027.items():
    apply_updates(by_id[uid], upd)
for uid, upd in TBA_RESOLVE.items():
    r = by_id[uid]
    apply_updates(r, upd)
    if "id" in upd:
        r["id"] = upd["id"]

# ----- NEW rows -----
def ev(id, name, acronym, typ, tags, sd, ed, loc, abs_dl, pap_dl, reg_dl, eb_dl,
       url, score, note, status, change, srcs):
    return {
        "id": id, "name": name, "acronym": acronym, "type": typ, "topic_tags": tags,
        "start_date": sd, "end_date": ed, "location": loc,
        "abstract_deadline": abs_dl, "paper_deadline": pap_dl,
        "registration_deadline": reg_dl, "early_bird_deadline": eb_dl,
        "website": url, "relevance_score": score, "relevance_note": note,
        "status": status, "last_verified": TODAY, "last_change": change, "source_urls": srcs,
    }

NEW_ROWS = [
    ev("cleo-2027", "Conference on Lasers and Electro-Optics 2027", "CLEO 2027", "Conference",
       "integrated quantum photonics;quantum optics;nonlinear optics;silicon photonics;single-photon sources",
       "2027-05-02", "2027-05-07", "Long Beach, CA, USA", "TBA", "n/a", "TBA", "TBA",
       "https://cleoconference.org/", 5,
       "Premier US laser/photonics conference; major IQP, Si-PIC and quantum optics sessions",
       "upcoming", "new entry; year roll-over from CLEO 2026; confirmed May 2-7 2027 Long Beach CA",
       "https://cleoconference.org/"),
    ev("qcrypt-2027", "QCrypt 2027", "QCrypt 2027", "Conference",
       "quantum communication;QKD;quantum networking",
       "TBA", "TBA", "TBA", "TBA", "n/a", "TBA", "TBA",
       "https://qcrypt.net/", 4,
       "Premier annual QKD/quantum-cryptography conference; host city bidding process underway, not yet publicly announced",
       "upcoming", "new entry; year roll-over placeholder, host TBA pending bid selection",
       "https://qcrypt.net/"),
    ev("ecio-2027", "European Conference on Integrated Optics 2027", "ECIO 2027", "Conference",
       "integrated quantum photonics;photonic integrated circuits;silicon photonics;nanophotonics",
       "TBA", "TBA", "TBA", "TBA", "TBA", "TBA", "TBA",
       "https://www.ecio-conference.org/", 5,
       "Largest European integrated optics conference (year roll-over); 2027 edition not yet announced",
       "upcoming", "new entry; placeholder, not yet announced as of 2026-08-16",
       "https://www.ecio-conference.org/"),
    ev("ieee-siphotonics-2027", "IEEE Silicon Photonics Conference 2027", "IEEE SiPho 2027", "Conference",
       "silicon photonics;photonic integrated circuits;III-V on Si",
       "TBA", "TBA", "TBA", "TBA", "TBA", "TBA", "TBA",
       "https://www.ieee-siphotonics.org/", 5,
       "Dedicated silicon photonics conference (year roll-over); 2027 edition not yet announced",
       "upcoming", "new entry; placeholder, not yet announced as of 2026-08-16",
       "https://www.ieee-siphotonics.org/"),
    ev("optica-quantum-2027", "Optica Quantum 2.0 Conference and Exhibition 2027", "Optica Quantum 2027", "Conference",
       "integrated quantum photonics;quantum optics;single-photon sources;quantum communication;quantum computing hardware",
       "2027-06-07", "2027-06-10", "Copenhagen, Denmark", "TBA", "n/a", "TBA", "TBA",
       "https://www.optica.org/events/topical_meetings/quantum/", 5,
       "Premier quantum photonics conference spanning PICs, detectors, QKD, and hardware (year roll-over)",
       "upcoming", "new entry; year roll-over, confirmed Jun 7-10 2027 Bella Center Copenhagen",
       "https://www.optica.org/events/topical_meetings/quantum/"),
    ev("damop-2027", "58th APS Division of Atomic, Molecular and Optical Physics Meeting", "DAMOP 2027", "Conference",
       "quantum optics;nonlinear optics;single-photon sources",
       "TBA", "TBA", "Chicago, IL, USA", "TBA", "TBA", "TBA", "TBA",
       "https://www.aps.org/events/2027/damop-2027", 3,
       "AMO physics flagship (year roll-over); venue confirmed Chicago, exact dates not yet published",
       "upcoming", "new entry; venue confirmed, dates TBA",
       "https://www.aps.org/events/2027/damop-2027"),
    ev("islc-2028", "31st International Semiconductor Laser Conference", "ISLC 2028", "Conference",
       "single-photon sources;III-V on Si;photonic integrated circuits",
       "TBA", "TBA", "TBA", "TBA", "TBA", "TBA", "TBA",
       "https://ieee-islc.org/", 4,
       "III-V semiconductor laser conference (biennial year roll-over); host not yet announced",
       "upcoming", "new entry; placeholder, biennial series, host TBA",
       "https://ieee-islc.org/"),
    ev("spie-pe-2028", "SPIE Photonics Europe 2028", "SPIE PE 2028", "Conference",
       "integrated quantum photonics;silicon photonics;nanophotonics",
       "2028-04-02", "2028-04-06", "Strasbourg, France", "TBA", "TBA", "TBA", "TBA",
       "https://spie.org/conferences-and-exhibitions/photonics-europe", 5,
       "Major European SPIE photonics conference (biennial year roll-over); confirmed",
       "upcoming", "new entry; confirmed Apr 2-6 2028 Strasbourg (beyond normal 18-month horizon but flagship biennial, worth tracking early)",
       "https://spie.org/conferences-and-exhibitions/photonics-europe"),
    ev("benasque-qnp-2027", "Quantum Nanophotonics 2027", "Benasque QNP 2027", "Workshop",
       "nanophotonics;quantum optics;single-photon sources",
       "2027-03-28", "2027-04-02", "Benasque, Spain", "TBA", "n/a", "TBA", "n/a",
       "https://www.benasque.org/2027quantumnanophotonics/", 5,
       "Directly on nanophotonics and single-photon sources (year roll-over)",
       "upcoming", "new entry; year roll-over from 2025/2026 pattern; dates per secondary source (Benasque site unreachable this run) - verify",
       "https://www.benasque.org/2027quantumnanophotonics/"),
    ev("grc-qclm-2027", "Gordon Research Conference — Quantum Control of Light and Matter 2027", "GRC QCLM 2027", "Conference",
       "quantum optics;quantum communication;quantum networking",
       "2027-08-01", "2027-08-06", "Newport, RI, USA", "n/a", "n/a", "TBA", "n/a",
       "https://www.grc.org/quantum-control-of-light-and-matter-conference/2027/", 4,
       "Core coherent-control/quantum-optics GRC, not previously tracked",
       "upcoming", "new entry",
       "https://www.grc.org/quantum-control-of-light-and-matter-conference/2027/"),
    ev("grc-quantum-sensing-2027", "Gordon Research Conference — Quantum Sensing 2027", "GRC QSens 2027", "Conference",
       "quantum optics;single-photon sources;SNSPDs",
       "2027-07-11", "2027-07-16", "Les Diablerets, Switzerland", "n/a", "n/a", "TBA", "n/a",
       "https://www.grc.org/quantum-sensing-conference/2027/", 3,
       "Photon-based sensing platforms overlap with AG Schuck's detector work",
       "upcoming", "new entry",
       "https://www.grc.org/quantum-sensing-conference/2027/"),
    ev("les-houches-nqio-2027", "Non-Linear, Quantum and Imaging Optics: from Fundamentals to Applications", "NQIO 2027", "Summer School",
       "quantum optics;nonlinear optics",
       "TBA", "TBA", "Les Houches, France", "n/a", "n/a", "TBA", "n/a",
       "https://www.houches-school-physics.com/program/program-2027/", 5,
       "New Les Houches doctoral school on nonlinear/quantum/imaging optics for 2027",
       "upcoming", "new entry; dates not yet published (applications open early 2027)",
       "https://www.houches-school-physics.com/program/program-2027/"),
    ev("les-houches-photai-2027", "Photonics & AI", "Photonics&AI 2027", "Summer School",
       "integrated quantum photonics;photonic integrated circuits",
       "TBA", "TBA", "Les Houches, France", "n/a", "n/a", "TBA", "n/a",
       "https://www.houches-school-physics.com/program/program-2027/", 4,
       "New Les Houches school on photonic accelerators, neuromorphic photonics and AI-assisted PIC design",
       "upcoming", "new entry; dates not yet published (applications open early 2027)",
       "https://www.houches-school-physics.com/program/program-2027/"),
    ev("ieee-qpain-2027", "IEEE 3rd International Conference on Quantum Photonics, AI & Networking", "IEEE QPAIN 2027", "Conference",
       "integrated quantum photonics;quantum networking",
       "2027-04-08", "2027-04-10", "Chattogram, Bangladesh", "TBA", "TBA", "TBA", "TBA",
       "https://qpain.org/", 3,
       "Quantum photonics and networking conference (year roll-over)",
       "upcoming", "new entry; year roll-over from 2026 edition",
       "https://qpain.org/"),
    ev("eqtc-2027", "European Quantum Technologies Conference 2027", "EQTC 2027", "Conference",
       "integrated quantum photonics;quantum communication;quantum computing hardware;quantum networking",
       "2027-06-22", "2027-06-25", "Munich, Germany", "TBA", "TBA", "TBA", "TBA",
       "https://qt.eu/events/", 5,
       "EU Quantum Flagship conference - all aspects of quantum technology, high visibility (year roll-over)",
       "upcoming", "new entry; year roll-over, venue Munich (unconfirmed - verify next run)",
       "https://qt.eu/events/"),
    ev("pd-workshop-2027", "8th International Workshop on New Photon-Detectors", "PD2027", "Workshop",
       "single-photon sources;SNSPDs",
       "TBA", "TBA", "Beijing, China", "TBA", "n/a", "TBA", "n/a",
       "https://web.infn.it/photondetectors2025/", 3,
       "Biennial HEP-instrumentation workshop covering SiPM/SNSPD/PMT detector technology",
       "upcoming", "new entry; dates not yet finalized (~May 2027 expected)",
       "https://web.infn.it/photondetectors2025/"),
    ev("iqn-hub-summerschool-2027", "UK Integrated Quantum Networks Hub Summer School 2027", "IQN Hub SS27", "Summer School",
       "quantum communication;quantum networking;single-photon sources;SNSPDs",
       "2027-06-27", "2027-07-03", "Edinburgh, UK", "n/a", "TBA", "n/a", "n/a",
       "https://iqnhub.org/2027-summer-school/", 4,
       "Residential UK quantum-networking school for PhD students/ECRs",
       "upcoming", "new entry",
       "https://iqnhub.org/2027-summer-school/"),
    ev("iqt-nordics-2027", "IQT Nordics 2027", "IQT Nordics 2027", "Conference",
       "quantum communication;quantum computing hardware;integrated quantum photonics",
       "2027-06-07", "2027-06-09", "Copenhagen, Denmark", "n/a", "n/a", "TBA", "TBA",
       "https://iqtevent.com/nordics/", 3,
       "Nordic quantum industry conference (year roll-over, venue moved Oslo to Copenhagen)",
       "upcoming", "new entry; year roll-over from 2026 Oslo edition, moved to Copenhagen",
       "https://iqtevent.com/nordics/"),
    ev("opic-2027", "Optics & Photonics International Congress 2027", "OPIC 2027", "Conference",
       "integrated quantum photonics;quantum optics;nanophotonics",
       "TBA", "TBA", "Yokohama, Japan", "TBA", "TBA", "TBA", "TBA",
       "https://opicon.jp/", 3,
       "Japanese international photonics congress with quantum tracks (year roll-over)",
       "upcoming", "new entry; placeholder for next edition after OPIC2026 (Apr 2026); not yet announced",
       "https://opicon.jp/"),
    ev("owtnm-2027", "33rd Optical Wave & Waveguide Theory and Modelling Workshop", "OWTNM 2027 (XXXIII)", "Workshop",
       "photonic integrated circuits;nanophotonics",
       "TBA", "TBA", "TBA", "TBA", "n/a", "TBA", "n/a",
       "https://www.owtnm.eu/", 3,
       "Waveguide modelling workshop relevant to PIC simulation (year roll-over)",
       "upcoming", "new entry; placeholder, host for XXXIII not yet announced",
       "https://www.owtnm.eu/"),
    ev("opon-next", "OPON Workshop — next edition (Optical Properties of Nanostructures)", "OPON next", "Workshop",
       "integrated quantum photonics;nanophotonics",
       "TBA", "TBA", "TBA", "TBA", "n/a", "TBA", "n/a",
       "TBA", 3,
       "Rotates Wrocław/Münster/Warsaw/Bayreuth; invitation-only; next host after Münster 2025 not yet announced",
       "upcoming", "new entry; placeholder for next OPON edition after Münster 2025",
       "original list"),
    ev("nrw-nano-2027", "12th NRW NanoConference 2027", "NRW Nano 2027", "Conference",
       "nanophotonics;photonic integrated circuits",
       "TBA", "TBA", "North Rhine-Westphalia, Germany", "TBA", "TBA", "TBA", "TBA",
       "https://www.nanoconference.de/", 3,
       "Regional NRW nanotechnology conference; local access (biennial year roll-over)",
       "upcoming", "new entry; placeholder for 12th edition after Dortmund 2025",
       "https://www.nanoconference.de/"),
    ev("icqe-next", "International Conference on Quantum Energy — next edition", "ICQE next", "Conference",
       "quantum computing hardware",
       "TBA", "TBA", "TBA", "TBA", "TBA", "TBA", "TBA",
       "https://icqe.com.au/", 2,
       "Quantum energy applications; adjacent to quantum computing hardware (year roll-over)",
       "upcoming", "new entry; placeholder, next edition after Padua 2025 not yet announced",
       "https://icqe.com.au/"),
]

all_rows = list(rows) + NEW_ROWS

# sanity: ensure all rows have all cols
for r in all_rows:
    for c in COLS:
        if c not in r or r[c] is None or r[c] == "":
            r[c] = r.get(c) or "TBA"

with open(OUT_TSV, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=COLS, delimiter="\t")
    w.writeheader()
    for r in all_rows:
        w.writerow({c: r[c] for c in COLS})

print(f"Wrote {len(all_rows)} rows to {OUT_TSV}")
print(f"Archived this run: {len(PASSED_IDS)}")
print(f"Updated (near-term): {len(NEAR_TERM)}")
print(f"Updated (2027): {len(FUTURE_2027)}")
print(f"Resolved TBA placeholders: {len(TBA_RESOLVE)}")
print(f"New rows added: {len(NEW_ROWS)}")


# ---- Part 2: build XLSX ----
import csv
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

TODAY = date(2026, 8, 16)
TSV_PATH = "/home/user/conferences/conferences_2026-08-16.tsv"
XLSX_PATH = "/home/user/conferences/conferences_2026-08-16.xlsx"

COLS = [
    "id", "name", "acronym", "type", "topic_tags",
    "start_date", "end_date", "location",
    "abstract_deadline", "paper_deadline",
    "registration_deadline", "early_bird_deadline",
    "website", "relevance_score", "relevance_note",
    "status", "last_verified", "last_change", "source_urls",
]
DEADLINE_COLS = ["abstract_deadline", "paper_deadline",
                 "registration_deadline", "early_bird_deadline"]

with open(TSV_PATH, newline="", encoding="utf-8") as f:
    events = list(csv.DictReader(f, delimiter="\t"))


def parse_date(s):
    if not s or s in ("TBA", "n/a"):
        return None
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None


def days_from_today(s):
    d = parse_date(s)
    if d is None:
        return None
    return (d - TODAY).days


RED = PatternFill(fill_type="solid", fgColor="FF4444")
AMBER = PatternFill(fill_type="solid", fgColor="FFAA00")


def deadline_fill(s):
    n = days_from_today(s)
    if n is None:
        return None
    if 0 <= n <= 14:
        return RED
    if 15 <= n <= 60:
        return AMBER
    return None


HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
ARCHIVED_FILL = PatternFill(fill_type="solid", fgColor="E8E8E8")
SCORE_COLORS = {5: "D4EDDA", 4: "D1ECF1", 3: "FFF3CD", 2: "F8D7DA", 1: "F5C6CB"}

wb = Workbook()
ws = wb.active
ws.title = "Conferences"

ws.append(COLS)
for cell in ws[1]:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws.row_dimensions[1].height = 28
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"

col_idx = {c: i + 1 for i, c in enumerate(COLS)}

for ev in events:
    row_vals = [ev.get(c, "") for c in COLS]
    ws.append(row_vals)
    row_num = ws.max_row

    is_archived = ev.get("status") == "archived"
    if is_archived:
        for cell in ws[row_num]:
            cell.fill = ARCHIVED_FILL

    score_cell = ws.cell(row=row_num, column=col_idx["relevance_score"])
    sc = ev.get("relevance_score")
    if sc and str(sc).isdigit() and int(sc) in SCORE_COLORS:
        score_cell.fill = PatternFill(fill_type="solid", fgColor=SCORE_COLORS[int(sc)])

    for dcol in DEADLINE_COLS:
        val = ev.get(dcol, "")
        fill = deadline_fill(str(val))
        if fill:
            ws.cell(row=row_num, column=col_idx[dcol]).fill = fill

CAP = {
    "id": 28, "name": 55, "acronym": 18, "type": 16, "topic_tags": 55,
    "start_date": 13, "end_date": 13, "location": 28,
    "abstract_deadline": 17, "paper_deadline": 14,
    "registration_deadline": 20, "early_bird_deadline": 18,
    "website": 48, "relevance_score": 9, "relevance_note": 52,
    "status": 14, "last_verified": 14, "last_change": 50, "source_urls": 60,
}
for col_name, cap_w in CAP.items():
    ci = col_idx[col_name]
    ws.column_dimensions[get_column_letter(ci)].width = cap_w

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=False, vertical="top")

wb.save(XLSX_PATH)
print(f"XLSX -> {XLSX_PATH} ({len(events)} rows)")

from collections import Counter
stats = Counter(e["status"] for e in events)
for s, n in sorted(stats.items()):
    print(f"  {s:15s}: {n}")
